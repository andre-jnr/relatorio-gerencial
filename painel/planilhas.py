"""Planilhas lidas pelo painel (relidas sozinhas quando são salvas):

- classificação dos pagamentos (PAGAMENTOS, CLASSIFICAÇÃO, GRUPO, PLANO DE CONTAS)
- vendas/estoque de unidade com "planilha_vendas" (CD): abas VENDAS e ESTOQUE.
  Pode ser trocada pelo painel (superusuário); a anterior vai para dados/backups/.
"""
import io
import logging
import os
import shutil
import threading
import time

from openpyxl import load_workbook

from .config import CFG, TODAS, caminho, normaliza

log = logging.getLogger("painel")
_trava = threading.Lock()

MESES_SIGLA = {"JAN": 1, "FEV": 2, "MAR": 3, "ABR": 4, "MAI": 5, "JUN": 6,
               "JUL": 7, "AGO": 8, "SET": 9, "OUT": 10, "NOV": 11, "DEZ": 12}


class ErroPlanilha(Exception):
    pass


def _ler_aba(arq, aba=None):
    """Linhas da aba (ou da primeira) como listas de valores. Lê em memória: funciona com o Excel aberto."""
    with open(arq, "rb") as f:
        conteudo = io.BytesIO(f.read())
    wb = load_workbook(conteudo, read_only=True, data_only=True)
    try:
        if aba is None:
            ws = wb.worksheets[0]
        else:
            nome = next((n for n in wb.sheetnames if normaliza(n) == normaliza(aba)), None)
            if nome is None:
                raise ErroPlanilha(f"A planilha {arq.name} não tem a aba '{aba}'.")
            ws = wb[nome]
        return [list(l) for l in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def _texto(v):
    return "" if v is None else str(v).strip()


def _numero(v):
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "."))
    except (TypeError, ValueError):
        return None


def _mes(v):
    if isinstance(v, (int, float)):
        return int(v)
    t = normaliza(v)
    if t.isdigit():
        return int(t)
    return MESES_SIGLA.get(t[:3], 0)


def _indice(cab, nome):
    return cab.index(nome) if nome in cab else -1


# ------------------------------------------------------------------
# Classificação dos pagamentos
# ------------------------------------------------------------------
_classificacao = {"mapa": {}, "contas": {}, "lido_em": None}


def _arquivo_classificacao():
    return caminho(CFG.get("planilha_classificacao") or "dados/classificação_pagamentos.xlsx")


def _atualizar_classificacao():
    arq = _arquivo_classificacao()
    if not arq.exists():
        raise ErroPlanilha("planilha de classificação dos pagamentos não encontrada")
    mod = arq.stat().st_mtime
    if mod == _classificacao["lido_em"]:
        return False
    linhas = _ler_aba(arq)
    cab = [normaliza(c) for c in linhas[0]] if linhas else []
    i_conta, i_classe = _indice(cab, "PAGAMENTOS"), _indice(cab, "CLASSIFICACAO")
    if i_conta < 0 or i_classe < 0:
        raise ErroPlanilha("a planilha de classificação precisa das colunas PAGAMENTOS e CLASSIFICAÇÃO na primeira linha")
    i_grupo, i_plano = _indice(cab, "GRUPO"), _indice(cab, "PLANO DE CONTAS")
    mapa, contas = {}, {}
    pega = lambda l, i: _texto(l[i]) if 0 <= i < len(l) else ""
    for l in linhas[1:]:
        conta, classe = normaliza(pega(l, i_conta)), normaliza(pega(l, i_classe))
        if not (conta and classe):
            continue
        mapa[conta] = classe
        contas[conta] = {
            "nome": " ".join(pega(l, i_conta).split()),
            "grupo": pega(l, i_grupo) or "SEM GRUPO",
            "plano": pega(l, i_plano) or "SEM PLANO DE CONTAS",
        }
    _classificacao.update(mapa=mapa, contas=contas, lido_em=mod)
    revenda = sum(1 for v in mapa.values() if v == "PRODUTOS PARA REVENDA")
    log.info("Planilha de classificação carregada: %d contas (%d de produtos para revenda)", len(mapa), revenda)
    return True


def classe_conta(conta):
    return _classificacao["mapa"].get(normaliza(conta))


def info_conta(conta):
    return _classificacao["contas"].get(normaliza(conta))


# ------------------------------------------------------------------
# Planilha de vendas (CD): aba VENDAS (MÊS, ANO, TIPO, CLIENTE, VLR FINAL) e ESTOQUE
# ------------------------------------------------------------------
_planilhas_vendas = {}  # caminho -> {"lido_em", "meses": {(ano, mes): {...}}, "estoque": {(ano, mes): valor}}


def _achar_cabecalho(linhas, exigidas):
    for i, l in enumerate(linhas):
        cab = [normaliza(c) for c in l]
        if all(e in cab for e in exigidas):
            return i, cab
    return -1, []


def _atualizar_planilha_vendas(u):
    arq = caminho(u["planilha_vendas"])
    if not arq.exists():
        raise ErroPlanilha(f"planilha de vendas de {u['nome']} não encontrada")
    mod = arq.stat().st_mtime
    atual = _planilhas_vendas.get(str(arq))
    if atual and atual["lido_em"] == mod:
        return False

    meses, estoque, ignoradas = _ler_planilha_vendas(u, arq)
    _planilhas_vendas[str(arq)] = {"lido_em": mod, "meses": meses, "estoque": estoque}
    log.info("Planilha de vendas de %s carregada: %d mês(es)%s", u["nome"], len(meses),
             f" ({ignoradas} linha(s) ignorada(s))" if ignoradas else "")
    return True


def _ler_planilha_vendas(u, arq):
    aba = u.get("aba_vendas") or "VENDAS"
    linhas = _ler_aba(arq, aba)
    i_cab, cab = _achar_cabecalho(linhas, ["TIPO", "VLR FINAL", "MES", "ANO"])
    if i_cab < 0:
        raise ErroPlanilha(f"a aba {aba} da planilha de {u['nome']} precisa das colunas MÊS, ANO, TIPO e VLR FINAL")
    i_mes, i_ano, i_tipo = cab.index("MES"), cab.index("ANO"), cab.index("TIPO")
    i_valor, i_cliente = cab.index("VLR FINAL"), _indice(cab, "CLIENTE")

    meses, ignoradas = {}, 0
    for l in linhas[i_cab + 1:]:
        if not any(_texto(v) for v in l):
            continue
        pega = lambda i: l[i] if 0 <= i < len(l) else None
        mes, ano, valor = _mes(pega(i_mes)), _numero(pega(i_ano)), _numero(pega(i_valor))
        if not (1 <= mes <= 12) or ano is None or valor is None:
            ignoradas += 1
            continue
        m = meses.setdefault((int(ano), mes), {"venda": 0.0, "transf": 0.0, "qtde": 0, "clientes": {}, "transf_clientes": {}})
        tipo = normaliza(pega(i_tipo))
        cliente = _texto(pega(i_cliente))
        if tipo.startswith("TRANSF"):
            m["transf"] += valor
            chave = normaliza(cliente)
            m["transf_clientes"][chave] = m["transf_clientes"].get(chave, 0.0) + valor
        elif tipo == "VENDA":
            m["venda"] += valor
            m["qtde"] += 1
            nome = cliente or "SEM CLIENTE"
            m["clientes"][nome] = m["clientes"].get(nome, 0.0) + valor
        else:
            ignoradas += 1

    return meses, _ler_estoque(u, arq), ignoradas


def substituir_planilha_vendas(u, conteudo):
    """Troca a planilha de vendas da unidade pela enviada, guardando a anterior em dados/backups/.

    A nova só entra se puder ser lida como a anterior (abas e colunas); senão nada muda.
    Devolve o resumo da planilha nova e o nome do backup."""
    arq = caminho(u["planilha_vendas"])
    temp = arq.with_name(f".{arq.stem}.enviada{arq.suffix}")
    with _trava:
        temp.write_bytes(conteudo)
        try:
            try:
                meses, estoque, _ = _ler_planilha_vendas(u, temp)
            except ErroPlanilha as e:  # a mensagem fala do arquivo temporário: troca por "enviada"
                raise ErroPlanilha(str(e).replace(f"A planilha {temp.name}", "A planilha enviada").rstrip(".")) from e
            except Exception as e:  # noqa: BLE001  arquivo que não é um .xlsx válido
                raise ErroPlanilha("o arquivo enviado não é uma planilha do Excel (.xlsx) válida") from e
            if not meses:
                raise ErroPlanilha(f"a aba {u.get('aba_vendas') or 'VENDAS'} do arquivo enviado não tem nenhuma linha válida")
            backup = None
            if arq.exists():
                pasta = caminho("dados/backups")
                pasta.mkdir(parents=True, exist_ok=True)
                backup = pasta / f"{arq.stem}_{time.strftime('%Y-%m-%d_%H%M%S')}{arq.suffix}"
                shutil.copy2(arq, backup)
            try:
                os.replace(temp, arq)  # troca de uma vez: quem estiver lendo nunca vê o arquivo pela metade
            except PermissionError as e:
                raise ErroPlanilha(f"{arq.name} está aberta em outro programa no servidor (Excel?). Feche e envie de novo") from e
        finally:
            temp.unlink(missing_ok=True)
    ultimo = max(meses)
    return {"meses": len(meses), "ultimo_mes": f"{ultimo[1]:02d}/{ultimo[0]}", "tem_estoque": bool(estoque),
            "backup": backup.name if backup else None}


def _ler_estoque(u, arq):
    aba = u.get("aba_estoque") or "ESTOQUE"
    try:
        linhas = _ler_aba(arq, aba)
    except ErroPlanilha as e:
        log.warning("%s O estoque de %s não será mostrado.", e, u["nome"])
        return {}
    i_cab, cab = _achar_cabecalho(linhas, ["VLR FINAL", "MES", "ANO"])
    if i_cab < 0:
        log.warning("A aba %s não tem MÊS, ANO e VLR FINAL; o estoque de %s não será mostrado.", aba, u["nome"])
        return {}
    i_mes, i_ano, i_valor = cab.index("MES"), cab.index("ANO"), cab.index("VLR FINAL")
    estoque = {}
    for l in linhas[i_cab + 1:]:
        pega = lambda i: l[i] if 0 <= i < len(l) else None
        mes, ano, valor = _mes(pega(i_mes)), _numero(pega(i_ano)), _numero(pega(i_valor))
        if 1 <= mes <= 12 and ano is not None and valor is not None:
            estoque[(int(ano), mes)] = valor  # posição do mês (se repetir, vale a última linha)
    return estoque


def mes_planilha(u, ano, mes):
    p = _planilhas_vendas.get(str(caminho(u["planilha_vendas"])))
    return p["meses"].get((ano, mes)) if p else None


def estoque_planilha(u, ano, mes):
    p = _planilhas_vendas.get(str(caminho(u["planilha_vendas"])))
    return p["estoque"].get((ano, mes)) if p else None


def tem_estoque(u):
    p = _planilhas_vendas.get(str(caminho(u["planilha_vendas"]))) if u.get("planilha_vendas") else None
    return bool(p and p["estoque"])


# ------------------------------------------------------------------
# Transferências do CD para a loja ("transferencias_cd"), lidas da planilha do CD
# ------------------------------------------------------------------
def unidade_cd(u):
    t = u.get("transferencias_cd")
    cd = TODAS.get(t["unidade"]) if t else None
    return cd if cd and cd.get("planilha_vendas") else None


def transferencias_cd(u, ano, mes):
    cd = unidade_cd(u)
    if not cd:
        return 0.0
    p = mes_planilha(cd, ano, mes)
    if not p:
        return 0.0
    return float(p["transf_clientes"].get(normaliza(u["transferencias_cd"]["cliente"]), 0.0))


def atualizar(lojas):
    """Relê as planilhas que mudaram. Devolve True se alguma mudou (o cache fica velho)."""
    with _trava:
        mudou = _atualizar_classificacao()
        com_planilha = {}
        for u in lojas:
            if u.get("planilha_vendas"):
                com_planilha[u["slug"]] = u
            cd = unidade_cd(u)
            if cd:
                com_planilha[cd["slug"]] = cd
        for u in com_planilha.values():
            mudou = _atualizar_planilha_vendas(u) or mudou
        return mudou
